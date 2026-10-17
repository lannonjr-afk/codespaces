"""Unused scratch copy of the attachment extraction logic from newai.py.

This file is intentionally kept separate and is not wired into the app.
"""

from __future__ import annotations

import base64
import os
import zipfile
from dataclasses import dataclass
from io import BytesIO
from typing import Optional

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from werkzeug.datastructures import FileStorage

MAX_DOCUMENT_CHARACTERS = 100_000
TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".json",
    ".py",
    ".pyw",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",
    ".html",
    ".css",
    ".xml",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".log",
    ".sql",
    ".sh",
    ".c",
    ".h",
    ".cpp",
    ".hpp",
    ".java",
    ".go",
    ".rs",
    ".rb",
    ".php",
    ".swift",
    ".kt",
    ".r",
    ".csv",
}
IMAGE_MIME_TYPES = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".webp": "image/webp",
}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".docx", ".pdf", ".xlsx"} | set(IMAGE_MIME_TYPES)


@dataclass
class ChatAttachment:
    filename: str
    text: Optional[str] = None
    data_url: Optional[str] = None


def _check_text_size(text: str) -> str:
    if len(text) > MAX_DOCUMENT_CHARACTERS:
        raise ValueError(
            "Extracted file text must be 100,000 characters or less per file."
        )
    if not text.strip():
        raise ValueError("No readable text was found in the selected file.")
    return text


def extract_attachment(uploaded_file: FileStorage) -> ChatAttachment:
    """Read one supported upload without persisting it."""
    filename = os.path.basename(uploaded_file.filename or "")
    extension = os.path.splitext(filename)[1].lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported file type. Supported extensions: {supported}.")

    content = uploaded_file.read()
    if not content:
        raise ValueError(f"{filename or 'The selected file'} is empty.")

    if extension in IMAGE_MIME_TYPES:
        mime_type = IMAGE_MIME_TYPES[extension]
        data_url = f"data:{mime_type};base64,{base64.b64encode(content).decode('ascii')}"
        return ChatAttachment(filename=filename, data_url=data_url)

    if extension in TEXT_EXTENSIONS:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise ValueError(f"{filename} must use UTF-8 encoding.") from error
        return ChatAttachment(filename=filename, text=_check_text_size(text))

    if extension == ".pdf":
        try:
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise ValueError("Password-protected PDFs are not supported.")
            pages = []
            character_count = 0
            for page in reader.pages:
                page_text = page.extract_text() or ""
                character_count += len(page_text)
                if character_count > MAX_DOCUMENT_CHARACTERS:
                    raise ValueError(
                        "Extracted file text must be 100,000 characters or less per file."
                    )
                pages.append(page_text)
            text = "\n".join(pages)
        except PdfReadError as error:
            raise ValueError(f"{filename} could not be read as a PDF.") from error
        return ChatAttachment(filename=filename, text=_check_text_size(text))

    if extension == ".xlsx":
        try:
            workbook = load_workbook(
                BytesIO(content), read_only=True, data_only=True
            )
        except (InvalidFileException, zipfile.BadZipFile) as error:
            raise ValueError(
                f"{filename} could not be read as an XLSX workbook."
            ) from error
        try:
            rows = []
            character_count = 0
            for worksheet in workbook.worksheets:
                rows.append(f"Sheet: {worksheet.title}")
                for row in worksheet.iter_rows(values_only=True):
                    row_text = "\t".join(
                        "" if value is None else str(value) for value in row
                    )
                    character_count += len(row_text)
                    if character_count > MAX_DOCUMENT_CHARACTERS:
                        raise ValueError(
                            "Extracted file text must be 100,000 characters or less per file."
                        )
                    if row_text.strip():
                        rows.append(row_text)
        except (InvalidFileException, ValueError, zipfile.BadZipFile) as error:
            raise ValueError(
                f"{filename} could not be read as an XLSX workbook."
            ) from error
        finally:
            workbook.close()
        text = "\n".join(rows)
        return ChatAttachment(filename=filename, text=_check_text_size(text))

    if extension == ".docx":
        try:
            document = Document(BytesIO(content))
        except (PackageNotFoundError, ValueError, zipfile.BadZipFile) as error:
            raise ValueError(
                f"{filename} could not be read as a Word document."
            ) from error

        parts = []
        character_count = 0
        for paragraph in document.paragraphs:
            paragraph_text = paragraph.text
            character_count += len(paragraph_text)
            if character_count > MAX_DOCUMENT_CHARACTERS:
                raise ValueError(
                    "Extracted file text must be 100,000 characters or less per file."
                )
            if paragraph_text:
                parts.append(paragraph_text)

        for table in document.tables:
            for row in table.rows:
                row_text = "\t".join(cell.text for cell in row.cells)
                character_count += len(row_text)
                if character_count > MAX_DOCUMENT_CHARACTERS:
                    raise ValueError(
                        "Extracted file text must be 100,000 characters or less per file."
                    )
                if row_text.strip():
                    parts.append(row_text)

        return ChatAttachment(
            filename=filename, text=_check_text_size("\n".join(parts))
        )

    raise ValueError(f"{filename} could not be read.")
