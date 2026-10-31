import base64
import json
import os
import re
import sqlite3
import textwrap
from dataclasses import dataclass
from datetime import datetime, timezone
from html import unescape
from io import BytesIO
from typing import Optional
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from flask import Flask, jsonify, render_template, request


OLLAMA_BASE_URL = os.getenv("OLLAMA_BASE_URL", "http://localhost:11434/v1")
GROQ_BASE_URL = "https://api.groq.com/openai/v1"
GROQ_API_KEY = os.getenv("GROQ_API_KEY")
GROQ_MODEL = os.getenv("GROQ_MODEL", "openai/gpt-oss-120b")
GROQ_VISION_MODEL = os.getenv("GROQ_VISION_MODEL")
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "llama3.1")

DB_PATH = os.getenv("DB_PATH", "agent_memory.db")
ENABLE_ATTACHMENTS = os.getenv("ENABLE_ATTACHMENTS", "false").strip().lower() in {
    "1",
    "true",
    "yes",
    "on",
}
CONTEXT_MESSAGES = 4
MAX_INPUT_CHARS = 200_000
MAX_HISTORY_CHARS = 75_000
MAX_CONTENT_LENGTH_BYTES = 16 * 1024 * 1024
MAX_DOCUMENT_CHARACTERS = 100_000
ATTACHMENT_CHUNK_CHARACTERS = 6_000

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".json",
    ".py",
    ".pyw",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",
    ".html",
    ".css",
    ".xml",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".log",
    ".sql",
    ".sh",
    ".c",
    ".h",
    ".cpp",
    ".hpp",
    ".java",
    ".go",
    ".rs",
    ".rb",
    ".php",
    ".swift",
    ".kt",
    ".r",
    ".csv",
}
IMAGE_MIME_TYPES = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".webp": "image/webp",
}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".docx", ".pdf", ".xlsx"} | set(IMAGE_MIME_TYPES)


class SimpleMemory:
    """Very small SQLite-backed memory store for user facts and recent chat history."""

    def __init__(self, db_path: str = DB_PATH):
        self.db_path = db_path
        self.conn = sqlite3.connect(self.db_path)
        self.conn.row_factory = sqlite3.Row
        self._create_tables()

    def _create_tables(self) -> None:
        self.conn.execute(
            """
            CREATE TABLE IF NOT EXISTS facts (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
        self.conn.execute(
            """
            CREATE TABLE IF NOT EXISTS history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                role TEXT NOT NULL,
                content TEXT NOT NULL,
                meta TEXT,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
        self.conn.commit()

    def set_fact(self, key: str, value: str) -> None:
        self.conn.execute(
            "INSERT INTO facts (key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value, updated_at=CURRENT_TIMESTAMP",
            (key, value),
        )
        self.conn.commit()

    def get_fact(self, key: str) -> Optional[str]:
        row = self.conn.execute("SELECT value FROM facts WHERE key = ?", (key,)).fetchone()
        return row["value"] if row else None

    def all_facts(self) -> dict[str, str]:
        rows = self.conn.execute("SELECT key, value FROM facts ORDER BY key").fetchall()
        return {row["key"]: row["value"] for row in rows}

    def delete_fact(self, key: str) -> bool:
        cursor = self.conn.execute("DELETE FROM facts WHERE key = ?", (key,))
        self.conn.commit()
        return cursor.rowcount > 0

    def add_history(self, role: str, content: str, meta: Optional[dict] = None) -> None:
        self.conn.execute(
            "INSERT INTO history (role, content, meta) VALUES (?, ?, ?)",
            (role, content, json.dumps(meta) if meta is not None else None),
        )
        self.conn.commit()

    def recent_history(self, limit: int = 10) -> list[dict[str, str]]:
        rows = self.conn.execute(
            "SELECT role, content FROM history ORDER BY id DESC LIMIT ?",
            (limit,),
        ).fetchall()
        items = [{"role": row["role"], "content": row["content"]} for row in rows]
        return list(reversed(items))

    def clear_history(self) -> int:
        cursor = self.conn.execute("DELETE FROM history")
        self.conn.commit()
        return cursor.rowcount

    def stats(self) -> dict[str, int]:
        fact_count = self.conn.execute("SELECT COUNT(*) FROM facts").fetchone()[0]
        history_count = self.conn.execute("SELECT COUNT(*) FROM history").fetchone()[0]
        return {"facts": fact_count, "history": history_count}

    def close(self) -> None:
        self.conn.close()


@dataclass
class ChatAttachment:
    filename: str
    text: Optional[str] = None
    data_url: Optional[str] = None


def _check_text_size(text: str) -> str:
    if len(text) > MAX_DOCUMENT_CHARACTERS:
        raise ValueError("Extracted file text must be 100,000 characters or less per file.")
    if not text.strip():
        raise ValueError("No readable text was found in the selected file.")
    return text


def extract_attachment(uploaded_file) -> ChatAttachment:
    """Read one supported upload without persisting it."""
    filename = os.path.basename(uploaded_file.filename or "")
    extension = os.path.splitext(filename)[1].lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported file type. Supported extensions: {supported}.")

    content = uploaded_file.read()
    if not content:
        raise ValueError(f"{filename or 'The selected file'} is empty.")

    if extension in IMAGE_MIME_TYPES:
        mime = IMAGE_MIME_TYPES[extension]
        data_url = f"data:{mime};base64,{base64.b64encode(content).decode('ascii')}"
        return ChatAttachment(filename=filename, data_url=data_url)

    if extension in TEXT_EXTENSIONS:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise ValueError(f"{filename} must use UTF-8 encoding.") from error
        return ChatAttachment(filename=filename, text=_check_text_size(text))

    if extension == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise ValueError("Password-protected PDFs are not supported.")
            pages = []
            for page in reader.pages:
                page_text = page.extract_text() or ""
                if page_text:
                    pages.append(page_text)
            text = "\n".join(pages)
            return ChatAttachment(filename=filename, text=_check_text_size(text))
        except Exception as error:  # pragma: no cover
            raise ValueError(f"{filename} could not be read as a PDF.") from error

    if extension == ".xlsx":
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            rows = []
            for worksheet in workbook.worksheets:
                rows.append(f"Sheet: {worksheet.title}")
                for row in worksheet.iter_rows(values_only=True):
                    row_text = "\t".join("" if value is None else str(value) for value in row)
                    if row_text.strip():
                        rows.append(row_text)
            text = "\n".join(rows)
            return ChatAttachment(filename=filename, text=_check_text_size(text))
        except Exception as error:  # pragma: no cover
            raise ValueError(f"{filename} could not be read as an XLSX workbook.") from error

    if extension == ".docx":
        try:
            from docx import Document

            document = Document(BytesIO(content))
            parts = []
            for paragraph in document.paragraphs:
                if paragraph.text:
                    parts.append(paragraph.text)
            for table in document.tables:
                for row in table.rows:
                    row_text = "\t".join(cell.text for cell in row.cells)
                    if row_text.strip():
                        parts.append(row_text)
            text = "\n".join(parts)
            return ChatAttachment(filename=filename, text=_check_text_size(text))
        except Exception as error:  # pragma: no cover
            raise ValueError(f"{filename} could not be read as a Word document.") from error

    raise ValueError(f"{filename} could not be read.")


def chunk_attachment_text(text: str, max_chars: int = ATTACHMENT_CHUNK_CHARACTERS) -> list[str]:
    if len(text) <= max_chars:
        return [text]
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + max_chars, len(text))
        if end < len(text):
            split = text.rfind("\n", start, end)
            if split > start:
                end = split
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        start = end
        if start < len(text) and text[start] == "\n":
            start += 1
    return chunks or [text]


def _extract_youtube_data(page_text: str):
    match = re.search(r"var ytInitialData = (\{.*?\});\s*</script>", page_text, re.S)
    if not match:
        return None
    try:
        return json.loads(match.group(1))
    except json.JSONDecodeError:
        return None


def _request_with_timeout(url: str, headers: Optional[dict[str, str]] = None):
    request = Request(url, headers=headers or {})
    return urlopen(request, timeout=20)


def lookup_youtube_subscribers(query: str) -> str:
    normalized = query.strip()
    if not normalized:
        return "Please provide a YouTube channel name or handle."

    handle = normalized
    if not handle.startswith("@") and "youtube.com" not in handle:
        handle = handle.strip("/")

    try:
        page_url = f"https://www.youtube.com/{handle}" if handle.startswith("@") else f"https://www.youtube.com/@{handle}" if not handle.startswith("http") else handle
        if handle.startswith("http"):
            page_url = handle
        response = _request_with_timeout(page_url)
        page_text = response.read().decode("utf-8", errors="replace")
    except (HTTPError, URLError, ValueError):
        return "Could not reach YouTube to look up that channel."

    data = _extract_youtube_data(page_text)
    if not data:
        return "Could not parse the YouTube channel page."

    title = None
    channel_id = None
    channel_url = None

    metadata = data.get("metadata", {}).get("channelMetadataRenderer", {})
    if metadata:
        title = metadata.get("title")
        channel_id = metadata.get("externalId")
        if channel_id:
            channel_url = f"https://www.youtube.com/channel/{channel_id}"

    if title is None:
        header = data.get("header", {}).get("c4TabbedHeaderRenderer", {})
        if header:
            title = header.get("title")

    if title is None:
        return "I couldn’t find a YouTube channel for that lookup."

    subs = None
    try:
        page = data.get("header", {}).get("pageHeaderRenderer", {})
        content = page.get("content", {})
        metadata_rows = content.get("pageHeaderViewModel", {}).get("metadata", {}).get("contentMetadataViewModel", {}).get("metadataRows", [])
        for row in metadata_rows:
            for part in row.get("metadataParts", []):
                label = part.get("text", {}).get("content") or part.get("text", {})
                if label:
                    subs = label
                    break
            if subs:
                break
    except (AttributeError, TypeError):
        subs = None

    if subs is None:
        return f"YouTube profile for {title} was found, but subscriber count could not be parsed."

    timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    if channel_url:
        return f"YouTube currently displays about {subs} subscribers for {title}. Channel: {channel_url}. Checked {timestamp}."
    return f"YouTube currently displays about {subs} subscribers for {title}. Checked {timestamp}."


class Aura:
    """Simple assistant wrapper that keeps a memory and returns helpful replies."""

    COMMAND_HELP = "Commands: /help, /memory, /clear, /stats, /remember key=value"

    def __init__(self, memory: Optional[SimpleMemory] = None):
        self.memory = memory or SimpleMemory(DB_PATH)

    def _facts_block(self) -> str:
        facts = self.memory.all_facts()
        if not facts:
            return "No remembered facts."
        return "\n".join(f"- {key}: {value}" for key, value in sorted(facts.items()))

    def _context_messages(self, user_input: str) -> list[dict[str, str]]:
        history = self.memory.recent_history(limit=CONTEXT_MESSAGES * 2)
        user_messages = [item for item in history if item["role"] == "user"]
        if user_input and user_input.strip():
            user_messages = [item for item in user_messages if item["content"].strip() != user_input.strip()]
        return user_messages[-CONTEXT_MESSAGES:]

    def _build_messages(self, user_input: str, attachments: Optional[list[ChatAttachment]] = None) -> list[dict[str, object]]:
        context = self._context_messages(user_input)
        messages: list[dict[str, object]] = [{"role": "system", "content": "You are AURA, a helpful chat assistant. Give concise, accurate answers."}]
        for item in context:
            messages.append({"role": item["role"], "content": item["content"]})

        content: list[dict[str, object]] = [{"type": "text", "text": user_input}]
        if attachments:
            for attachment in attachments:
                if attachment.data_url:
                    content.append({"type": "image_url", "image_url": {"url": attachment.data_url}})
                elif attachment.text:
                    content.append({"type": "text", "text": f"Attachment: {attachment.filename}\n\n{attachment.text}"})
        messages.append({"role": "user", "content": content})
        return messages

    def _complete(self, messages: list[dict[str, object]], model: str) -> str:
        payload = {
            "model": model,
            "messages": messages,
            "temperature": 0.2,
            "max_tokens": 800,
        }
        body = json.dumps(payload).encode("utf-8")

        if GROQ_API_KEY:
            headers = {
                "Authorization": f"Bearer {GROQ_API_KEY}",
                "Content-Type": "application/json",
            }
            url = f"{GROQ_BASE_URL}/chat/completions"
        else:
            headers = {"Content-Type": "application/json"}
            url = f"{OLLAMA_BASE_URL.rstrip('/')}/chat/completions"

        request = Request(url, data=body, headers=headers, method="POST")
        with urlopen(request, timeout=60) as response:
            data = json.loads(response.read().decode("utf-8"))

        if "choices" not in data or not data["choices"]:
            raise RuntimeError("The model API returned no completion choices.")

        message = data["choices"][0].get("message", {})
        content = message.get("content") or message.get("text") or ""
        return str(content).strip()

    def answer_directly(self, user_input: str, attachments: Optional[list[ChatAttachment]] = None) -> str:
        text = (user_input or "").strip()
        if not text:
            return "Please provide a message."

        lowered = text.lower()
        if lowered in {"/help", "help"}:
            return self.COMMAND_HELP
        if lowered in {"/memory", "memory"}:
            return self._facts_block()
        if lowered in {"/stats", "stats"}:
            stats = self.memory.stats()
            return f"Facts: {stats['facts']} | History: {stats['history']}"
        if lowered in {"/clear", "clear"}:
            self.memory.clear_history()
            return "Conversation history cleared."
        if lowered.startswith("/remember "):
            _, _, payload = text.partition(" ")
            if not payload or "=" not in payload:
                return "Usage: /remember key=value"
            key, value = payload.split("=", 1)
            self.memory.set_fact(key.strip(), value.strip())
            return f"Remembered: {key.strip()} = {value.strip()}"

        if re.search(r"(?:how many|subscriber|subscribers|subs)\b.*(?:@|youtube|channel)", lowered):
            match = re.search(r"@?[A-Za-z0-9_.-]+", text)
            lookup_name = match.group(0) if match else text
            return lookup_youtube_subscribers(lookup_name)

        if attachments:
            if any(att.data_url for att in attachments):
                model = GROQ_VISION_MODEL or (GROQ_MODEL if GROQ_API_KEY else OLLAMA_MODEL)
            else:
                model = GROQ_MODEL if GROQ_API_KEY else OLLAMA_MODEL
            messages = self._build_messages(text, attachments)
            return self._complete(messages, model)

        try:
            messages = self._build_messages(text)
            return self._complete(messages, GROQ_MODEL if GROQ_API_KEY else OLLAMA_MODEL)
        except Exception:
            recent = self._context_messages(text)
            if not recent:
                return "I’m ready to help."
            summary = "\n".join(f"{item['role']}: {item['content']}" for item in recent)
            return (
                "I can help with that. "
                f"Recent chat:\n{summary}\n\n"
                "If you want, ask me for a summary, a direct answer, or a rewrite."
            )

    def handle(self, user_input: str) -> str:
        return self.answer_directly(user_input)

    def answer_and_record(self, user_input: str, attachments: Optional[list[ChatAttachment]] = None) -> str:
        answer = self.answer_directly(user_input, attachments)
        self.memory.add_history("user", user_input.strip())
        self.memory.add_history("assistant", answer)
        return answer

    def _record(self, user_input: str, answer: str, meta: Optional[dict] = None) -> None:
        self.memory.add_history("user", user_input.strip(), meta)
        self.memory.add_history("assistant", answer, meta)

    def run_command(self, text: str) -> bool:
        text = text.strip()
        if not text:
            return False
        if text.startswith("/"):
            self.answer_and_record(text)
            return True
        return False


app = Flask(__name__)
aura = Aura(memory=SimpleMemory(DB_PATH))


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/healthz")
def healthz():
    return jsonify({"status": "ok"})


@app.route("/api/chat", methods=["POST"])
def chat():
    payload = request.get_json(silent=True) or {}
    user_input = str(payload.get("message", "")).strip()
    if not user_input:
        return jsonify({"error": "Message is required."}), 400

    attachments = []
    if ENABLE_ATTACHMENTS:
        for key in ("document", "documents"):
            files = request.files.getlist(key)
            for file in files:
                if not file or not file.filename:
                    continue
                try:
                    attachments.append(extract_attachment(file))
                except ValueError as error:
                    return jsonify({"error": str(error)}), 400

    if attachments:
        for attachment in attachments:
            if attachment.text:
                text_chunks = chunk_attachment_text(attachment.text)
                if len(text_chunks) > 1:
                    summary_parts = []
                    for index, chunk in enumerate(text_chunks):
                        summary_parts.append(f"Chunk {index + 1}:\n{chunk}")
                    attachment.text = "\n\n".join(summary_parts)

    response = aura.answer_and_record(user_input, attachments or None)
    return jsonify({"response": response})


def main(db_path: str = DB_PATH) -> None:
    global aura
    aura = Aura(memory=SimpleMemory(db_path))
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "8000")), debug=False)


if __name__ == "__main__":
    main()
