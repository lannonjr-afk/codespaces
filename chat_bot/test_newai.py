from io import BytesIO
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import newai
from docx import Document
from openpyxl import Workbook


class WebAppTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.db_path_patch = patch.object(
            newai, "DB_PATH", str(Path(self.temp_dir.name) / "test.db")
        )
        self.db_path_patch.start()
        self.client = newai.app.test_client()

    def tearDown(self):
        self.db_path_patch.stop()
        self.temp_dir.cleanup()

    def test_health_check_is_public(self):
        response = self.client.get("/healthz")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json, {"status": "ok"})

    def test_chat_page_is_public(self):
        response = self.client.get("/")

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"AURA", response.data)

    def test_chat_page_and_api(self):
        with patch.object(newai.Aura, "answer_directly", return_value="Hello"):
            response = self.client.post("/api/chat", json={"message": "Hi"})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json, {"response": "Hello"})
        memory = newai.SimpleMemory(newai.DB_PATH)
        try:
            self.assertEqual(
                memory.recent_history(),
                [
                    {"role": "user", "content": "Hi"},
                    {"role": "assistant", "content": "Hello"},
                ],
            )
        finally:
            memory.close()

    def test_chat_rejects_empty_message(self):
        response = self.client.post("/api/chat", json={"message": "  "})

        self.assertEqual(response.status_code, 400)

    def test_txt_upload_is_used_but_not_saved_to_history(self):
        with patch.object(
            newai.Aura, "answer_directly", return_value="It says hello."
        ) as answer_directly:
            response = self.client.post(
                "/api/chat",
                data={
                    "message": "What does this say?",
                    "document": (BytesIO(b"Hello from a file."), "note.txt"),
                },
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json, {"response": "It says hello."})
        answer_directly.assert_called_once_with(
            "What does this say?",
            [newai.ChatAttachment("note.txt", text="Hello from a file.")],
        )
        memory = newai.SimpleMemory(newai.DB_PATH)
        try:
            self.assertEqual(
                memory.recent_history(),
                [
                    {"role": "user", "content": "What does this say?"},
                    {"role": "assistant", "content": "It says hello."},
                ],
            )
        finally:
            memory.close()

    def test_attachment_chunks_cover_full_text_with_overlap(self):
        text = "".join(f"line {index:04d}\n" for index in range(1500))
        chunks = newai.chunk_attachment_text(text)

        self.assertGreater(len(chunks), 1)
        self.assertTrue(
            all(
                len(chunk) <= newai.ATTACHMENT_CHUNK_CHARACTERS
                for chunk in chunks
            )
        )

        reconstructed = chunks[0]
        for chunk in chunks[1:]:
            overlap = min(len(reconstructed), len(chunk))
            while overlap and not reconstructed.endswith(chunk[:overlap]):
                overlap -= 1
            self.assertGreater(overlap, 0)
            reconstructed += chunk[overlap:]

        self.assertEqual(reconstructed, text)

    def test_long_attachment_is_analyzed_in_chunks_before_final_answer(self):
        memory = newai.SimpleMemory(str(Path(self.temp_dir.name) / "chunks.db"))
        text = (
            "START OF FILE\n"
            + ("A" * newai.ATTACHMENT_CHUNK_CHARACTERS)
            + "\nMIDDLE OF FILE\n"
            + ("B" * newai.ATTACHMENT_CHUNK_CHARACTERS)
            + "\nEND OF FILE"
        )
        chunks = newai.chunk_attachment_text(text)
        notes = [f"segment {index} notes" for index in range(len(chunks))]
        try:
            aura = newai.Aura(memory)
            with patch.object(
                aura,
                "_complete",
                side_effect=[*notes, "Final answer"],
            ) as complete:
                answer = aura.answer_directly(
                    "Explain this file.",
                    [newai.ChatAttachment("new_ai.py", text=text)],
                )

            self.assertEqual(answer, "Final answer")
            self.assertEqual(complete.call_count, len(chunks) + 1)
            segment_requests = [
                call.args[0][-1]["content"]
                for call in complete.call_args_list[:-1]
            ]
            self.assertIn("START OF FILE", segment_requests[0])
            self.assertIn("MIDDLE OF FILE", "\n".join(segment_requests))
            self.assertIn("END OF FILE", segment_requests[-1])
            final_request = complete.call_args_list[-1].args[0][-1]["content"]
            for note in notes:
                self.assertIn(note, final_request[1]["text"])
        finally:
            memory.close()

    def test_pdf_upload_extracts_text(self):
        page = Mock()
        page.extract_text.return_value = "PDF document text."
        reader = Mock(is_encrypted=False, pages=[page])

        with (
            patch.object(newai, "PdfReader", return_value=reader),
            patch.object(newai.Aura, "answer_directly", return_value="Answer."),
        ):
            response = self.client.post(
                "/api/chat",
                data={
                    "message": "Summarize this.",
                    "document": (BytesIO(b"%PDF-test"), "document.pdf"),
                },
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json, {"response": "Answer."})

    def test_upload_rejects_unsupported_files(self):
        response = self.client.post(
            "/api/chat",
            data={
                "message": "Read this.",
                "document": (BytesIO(b"data"), "document.unsupported"),
            },
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("Unsupported file type", response.json["error"])

    def test_python_file_is_read_as_text(self):
        with patch.object(
            newai.Aura, "answer_directly", return_value="This prints hello."
        ) as answer_directly:
            response = self.client.post(
                "/api/chat",
                data={
                    "message": "What does this code do?",
                    "documents": (BytesIO(b'print("hello")'), "example.py"),
                },
            )

        self.assertEqual(response.status_code, 200)
        answer_directly.assert_called_once_with(
            "What does this code do?",
            [newai.ChatAttachment("example.py", text='print("hello")')],
        )

    def test_xlsx_upload_extracts_cells(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Sales"
        worksheet.append(["Product", "Units"])
        worksheet.append(["Apples", 12])
        content = BytesIO()
        workbook.save(content)

        with patch.object(
            newai.Aura, "answer_directly", return_value="There are 12 apples."
        ) as answer_directly:
            response = self.client.post(
                "/api/chat",
                data={
                    "message": "How many apples?",
                    "documents": (BytesIO(content.getvalue()), "sales.xlsx"),
                },
            )

        self.assertEqual(response.status_code, 200)
        attachment = answer_directly.call_args.args[1][0]
        self.assertEqual(attachment.filename, "sales.xlsx")
        self.assertIn("Sheet: Sales", attachment.text)
        self.assertIn("Apples\t12", attachment.text)

    def test_docx_upload_extracts_paragraphs_and_tables(self):
        document = Document()
        document.add_paragraph("Quarterly report")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Revenue"
        table.cell(0, 1).text = "$100"
        content = BytesIO()
        document.save(content)

        with patch.object(
            newai.Aura, "answer_directly", return_value="Revenue is $100."
        ) as answer_directly:
            response = self.client.post(
                "/api/chat",
                data={
                    "message": "What is the revenue?",
                    "documents": (BytesIO(content.getvalue()), "report.docx"),
                },
            )

        self.assertEqual(response.status_code, 200)
        attachment = answer_directly.call_args.args[1][0]
        self.assertEqual(attachment.filename, "report.docx")
        self.assertIn("Quarterly report", attachment.text)
        self.assertIn("Revenue\t$100", attachment.text)

    def test_image_upload_requires_vision_model_setting(self):
        response = self.client.post(
            "/api/chat",
            data={
                "message": "Describe this.",
                "documents": (BytesIO(b"image-bytes"), "photo.png"),
            },
        )

        self.assertEqual(response.status_code, 503)
        self.assertIn("Set GROQ_VISION_MODEL", response.json["error"])

    def test_image_upload_uses_configured_vision_model(self):
        with (
            patch.object(newai, "GROQ_VISION_MODEL", "vision-model"),
            patch.object(newai.Aura, "answer_directly", return_value="A photo.")
            as answer_directly,
        ):
            response = self.client.post(
                "/api/chat",
                data={
                    "message": "Describe this.",
                    "documents": (BytesIO(b"image-bytes"), "photo.png"),
                },
            )

        self.assertEqual(response.status_code, 200)
        attachment = answer_directly.call_args.args[1][0]
        self.assertEqual(attachment.filename, "photo.png")
        self.assertTrue(attachment.data_url.startswith("data:image/png;base64,"))

    def test_image_and_text_attachments_use_multimodal_vision_request(self):
        memory = newai.SimpleMemory(str(Path(self.temp_dir.name) / "vision.db"))
        try:
            aura = newai.Aura(memory)
            with (
                patch.object(newai, "GROQ_VISION_MODEL", "vision-model"),
                patch.object(aura, "_complete", return_value="Answer.") as complete,
            ):
                answer = aura.answer_directly(
                    "Explain these files.",
                    [
                        newai.ChatAttachment("script.py", text="print('hello')"),
                        newai.ChatAttachment(
                            "photo.png", data_url="data:image/png;base64,aGVsbG8="
                        ),
                    ],
                )

            self.assertEqual(answer, "Answer.")
            messages = complete.call_args.args[0]
            user_content = messages[-1]["content"]
            self.assertIn("script.py", user_content[1]["text"])
            self.assertEqual(user_content[2]["type"], "image_url")
            self.assertEqual(complete.call_args.kwargs["model"], "vision-model")
        finally:
            memory.close()

    def test_upload_rejects_non_utf8_text(self):
        response = self.client.post(
            "/api/chat",
            data={
                "message": "Read this.",
                "document": (BytesIO(b"\xff"), "document.txt"),
            },
        )

        self.assertEqual(response.status_code, 400)
        self.assertEqual(
            response.json, {"error": "document.txt must use UTF-8 encoding."}
        )

    def test_upload_rejects_extracted_text_over_limit(self):
        response = self.client.post(
            "/api/chat",
            data={
                "message": "Read this.",
                "document": (
                    BytesIO(b"a" * (newai.MAX_DOCUMENT_CHARACTERS + 1)),
                    "document.txt",
                ),
            },
        )

        self.assertEqual(response.status_code, 400)
        self.assertEqual(
            response.json,
            {
                "error": (
                    "Extracted file text must be 100,000 characters or less "
                    "per file."
                )
            },
        )


class YouTubeLookupTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.db_path_patch = patch.object(
            newai, "DB_PATH", str(Path(self.temp_dir.name) / "test.db")
        )
        self.db_path_patch.start()
        self.client = newai.app.test_client()

    def tearDown(self):
        self.db_path_patch.stop()
        self.temp_dir.cleanup()

    @staticmethod
    def _channel_page(title="Example Creator", count="1.23M subscribers"):
        import json

        data = {
            "header": {
                "pageHeaderRenderer": {
                    "content": {
                        "pageHeaderViewModel": {
                            "metadata": {
                                "contentMetadataViewModel": {
                                    "metadataRows": [
                                        {
                                            "metadataParts": [
                                                {
                                                    "text": {"content": count},
                                                    "accessibilityLabel": count,
                                                }
                                            ]
                                        }
                                    ]
                                }
                            }
                        }
                    }
                }
            },
            "metadata": {
                "channelMetadataRenderer": {
                    "title": title,
                    "externalId": "UC123",
                }
            },
        }
        return f"var ytInitialData = {json.dumps(data)};"

    @staticmethod
    def _search_page(*channels):
        import json

        data = {
            "contents": [
                {
                    "channelRenderer": {
                        "channelId": channel_id,
                        "title": {"simpleText": title},
                    }
                }
                for title, channel_id in channels
            ]
        }
        return f"var ytInitialData = {json.dumps(data)};"

    def test_handle_lookup_returns_subscriber_count_source_and_check_time(self):
        response = BytesIO(self._channel_page().encode())
        with patch.object(newai, "urlopen", return_value=response) as urlopen:
            answer = newai.lookup_youtube_subscribers("@example")

        self.assertIn(
            "YouTube currently displays about 1.23M subscribers for "
            "Example Creator.",
            answer,
        )
        self.assertIn("https://www.youtube.com/channel/UC123", answer)
        self.assertRegex(answer, r"checked \d{4}-\d{2}-\d{2} \d{2}:\d{2} UTC")
        self.assertEqual(
            urlopen.call_args.args[0].full_url, "https://www.youtube.com/@example"
        )

    def test_name_search_asks_user_to_disambiguate_instead_of_guessing(self):
        response = BytesIO(
            self._search_page(
                ("Demo Creator", "UC111"), ("Demo Gaming", "UC222")
            ).encode()
        )
        with patch.object(newai, "urlopen", return_value=response) as urlopen:
            answer = newai.lookup_youtube_subscribers("Demo")

        self.assertIn("couldn't reliably identify the right one", answer)
        self.assertIn("Demo Creator", answer)
        self.assertIn("Demo Gaming", answer)
        urlopen.assert_called_once()

    def test_exact_name_search_fetches_channel_page(self):
        search_result = BytesIO(
            self._search_page(
                ("Example Creator", "UC123"),
                ("Example  Creator", "UC456"),
            ).encode()
        )
        channel_page = BytesIO(self._channel_page().encode())
        with patch.object(
            newai, "urlopen", side_effect=[search_result, channel_page]
        ) as urlopen:
            answer = newai.lookup_youtube_subscribers("Example Creator")

        self.assertIn("1.23M subscribers", answer)
        self.assertEqual(urlopen.call_count, 2)
        self.assertEqual(
            urlopen.call_args_list[1].args[0].full_url,
            "https://www.youtube.com/channel/UC123",
        )

    def test_exact_spelling_takes_precedence_over_similar_name(self):
        search_result = BytesIO(
            self._search_page(
                ("MrBeast", "UC123"),
                ("Mr beast", "UC456"),
            ).encode()
        )
        channel_page = BytesIO(self._channel_page(title="MrBeast").encode())
        with patch.object(
            newai, "urlopen", side_effect=[search_result, channel_page]
        ) as urlopen:
            answer = newai.lookup_youtube_subscribers("MrBeast")

        self.assertIn("MrBeast", answer)
        self.assertIn("UC123", urlopen.call_args_list[1].args[0].full_url)

    def test_subscriber_question_uses_live_lookup_without_calling_model(self):
        memory = newai.SimpleMemory(":memory:")
        response = BytesIO(self._channel_page().encode())
        try:
            aura = newai.Aura(memory)
            with (
                patch.object(newai, "urlopen", return_value=response),
                patch.object(aura, "_complete") as complete,
            ):
                answer = aura.answer_directly(
                    "How many subs does @example have?"
                )

            self.assertIn(
                "YouTube currently displays about 1.23M subscribers for "
                "Example Creator.",
                answer,
            )
            complete.assert_not_called()
        finally:
            memory.close()

    def test_possessive_subscriber_question_is_recognized(self):
        memory = newai.SimpleMemory(":memory:")
        response = BytesIO(self._channel_page().encode())
        try:
            aura = newai.Aura(memory)
            with (
                patch.object(newai, "urlopen", return_value=response),
                patch.object(aura, "_complete") as complete,
            ):
                answer = aura.answer_directly(
                    "What's @example's subscriber count?"
                )

            self.assertIn("Example Creator", answer)
            complete.assert_not_called()
        finally:
            memory.close()

    def test_network_errors_are_reported_by_chat_endpoint(self):
        with patch.object(
            newai, "urlopen", side_effect=newai.URLError("offline")
        ):
            response = self.client.post(
                "/api/chat",
                json={"message": "How many subscribers does @example have?"},
            )

        self.assertEqual(response.status_code, 502)
        self.assertIn("Could not reach YouTube", response.json["error"])


if __name__ == "__main__":
    unittest.main()
